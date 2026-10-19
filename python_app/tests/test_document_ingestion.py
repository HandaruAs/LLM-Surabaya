import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from python_app.main import read_file


class DocumentExtractionTests(unittest.TestCase):
    def test_extracts_text_from_xlsx_workbook_sheets(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "anggaran.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Pendapatan"
            sheet.append(["Kode", "Uraian", "Anggaran"])
            sheet.append(["4.1", "Pajak daerah", 125000])
            workbook.create_sheet("Belanja").append(["Belanja modal", 65000])
            workbook.save(path)

            text = read_file(path)

        self.assertIn("[Sheet: Pendapatan]", text)
        self.assertIn("Pajak daerah", text)
        self.assertIn("125000", text)
        self.assertIn("[Sheet: Belanja]", text)
        self.assertIn("Belanja modal", text)

    def test_reports_empty_workbook_as_no_extractable_text(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "kosong.xlsx"
            Workbook().save(path)
            self.assertEqual(read_file(path), "")
